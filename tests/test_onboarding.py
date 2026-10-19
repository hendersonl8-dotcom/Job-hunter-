import io
import json
from types import SimpleNamespace

import pytest
from docx import Document
from fpdf import FPDF

from job_hunter import onboarding
from job_hunter.resume import base_resume
from job_hunter.scoring import Profile

RESUME_TEXT = """Jane Doe
jane@example.com | 206-555-0100 | Seattle, WA | Active Secret clearance
Help Desk Lead, Acme Corp, 2021-2025
- Resolved 40 tickets a week and trained 5 technicians on Active Directory and PowerShell.
BA Information Systems, 2020
"""


def facts(**over):
    data = dict(
        contact=onboarding.Contact(name="Jane Doe", email="jane@example.com", phone="206-555-0100",
                                   location="Seattle, WA", clearance="Active Secret"),
        tracks=[onboarding.Track(name="IT Support", title_words=["help desk", "it support"],
                                 headline="IT Support Lead", summary="Led 5 technicians.")],
        skills=[onboarding.Skill(name="Active Directory", group="Systems", aliases=["AD DS"]),
                onboarding.Skill(name="PowerShell", group="Systems")],
        experience=[onboarding.Role(title="Help Desk Lead", org="Acme Corp", dates="2021-2025",
                                    bullets=["Resolved 40 tickets a week and trained 5 technicians."])],
        education=[onboarding.Education(name="BA Information Systems", dates="2020")],
    )
    data.update(over)
    return onboarding.ExtractedResume(**data)


def test_extract_text_docx():
    doc = Document()
    for line in RESUME_TEXT.splitlines():
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    assert "Help Desk Lead" in onboarding.extract_text("r.docx", buf.getvalue())


def test_extract_text_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in RESUME_TEXT.splitlines():
        pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    assert "Help Desk Lead" in onboarding.extract_text("r.pdf", bytes(pdf.output()))


def test_extract_text_rejects_other_files_and_empty():
    with pytest.raises(ValueError):
        onboarding.extract_text("r.txt", b"x")
    doc = Document()
    buf = io.BytesIO()
    doc.save(buf)
    with pytest.raises(ValueError):
        onboarding.extract_text("r.docx", buf.getvalue())


def test_verify_keeps_supported_facts():
    out, warnings = onboarding.verify(facts(), RESUME_TEXT)
    assert warnings == []
    assert [s.name for s in out.skills] == ["Active Directory", "PowerShell"]


def test_verify_drops_invented_things():
    bad = facts(
        skills=[onboarding.Skill(name="Terraform", group="Cloud"), onboarding.Skill(name="PowerShell", group="Systems")],
        contact=onboarding.Contact(name="Jane Doe", email="other@example.com", phone="999-999-9999", clearance="TS/SCI"),
        experience=[onboarding.Role(title="Help Desk Lead", org="Acme Corp", bullets=["Saved $900,000 a year."])],
    )
    out, warnings = onboarding.verify(bad, RESUME_TEXT)
    assert [s.name for s in out.skills] == ["PowerShell"]
    assert out.contact.email == out.contact.phone == out.contact.clearance == ""
    joined = " ".join(warnings)
    assert "Terraform" in joined and "900,000" in joined and "TS/SCI" in joined


def test_built_files_load_in_the_app(tmp_path):
    f = facts()
    resume, profile = onboarding.build_resume(f), onboarding.build_profile(f, 65000, ["Seattle"])
    onboarding.save_json(tmp_path / "config" / "resume.json", resume)
    onboarding.save_json(tmp_path / "config" / "profile.json", profile)
    loaded = Profile.load(tmp_path / "config" / "profile.json")
    assert loaded.min_salary == 65000 and loaded.clearance == "Active Secret"
    assert loaded.skills["Active Directory"][1] is True
    assert "IT Support" in loaded.tracks
    base = base_resume(onboarding.load_json(tmp_path / "config" / "resume.json"), "IT Support")
    assert base["headline"] == "IT Support Lead"
    assert base["skills"] == [{"name": "Systems", "items": ["Active Directory", "PowerShell"]}]
    assert base_resume(resume, "Unknown")["contact"]["name"] == "Jane Doe"  # falls back to "Other"


def test_needs_onboarding(tmp_path):
    from job_hunter.config import ROOT
    assert onboarding.needs_onboarding(tmp_path / "p.json", tmp_path / "r.json")
    assert not onboarding.needs_onboarding(ROOT / "config" / "profile.json", ROOT / "config" / "resume.json")


def test_extract_uses_client_without_a_key():
    parsed = facts()
    seen = {}

    def parse(**kw):
        seen.update(kw)
        return SimpleNamespace(stop_reason="end_turn", parsed_output=parsed)

    client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(parse=parse)))
    assert onboarding.extract(RESUME_TEXT, client=client) is parsed
    assert "Help Desk Lead" in seen["messages"][0]["content"]


def test_extract_reports_refusal_and_truncation():
    for reason, parsed in (("refusal", None), ("max_tokens", None)):
        client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(
            parse=lambda **kw: SimpleNamespace(stop_reason=reason, parsed_output=parsed))))
        with pytest.raises(RuntimeError):
            onboarding.extract("x" * 100, client=client)
