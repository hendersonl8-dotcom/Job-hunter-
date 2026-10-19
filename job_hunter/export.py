"""Write a resume dict to Word (.docx) and PDF."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF


def contact_line(contact: dict) -> str:
    keys = ("location", "phone", "email", "linkedin", "github", "clearance")
    return "  |  ".join(contact[k] for k in keys if contact.get(k))


def to_docx(resume: dict) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(50)
        section.top_margin = section.bottom_margin = Pt(40)

    def para(text="", bold=False, size=None, align=None, space_after=2):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def heading(text):
        p = para(text.upper(), bold=True, size=11.5, space_after=3)
        p.paragraph_format.space_before = Pt(8)

    def bullet(text):
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)

    contact = resume["contact"]
    para(contact["name"], bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(resume["headline"], size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(contact_line(contact), size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    heading("Summary")
    para(resume["summary"])
    heading("Skills")
    for group in resume["skills"]:
        p = doc.add_paragraph()
        p.add_run(f"{group['name']}: ").bold = True
        p.add_run(", ".join(group["items"]))
        p.paragraph_format.space_after = Pt(1)
    heading("Experience")
    for role in resume["experience"]:
        p = doc.add_paragraph()
        p.add_run(role["title"]).bold = True
        p.add_run(f", {role['org']}")
        p.add_run(f"    {role['dates']}").italic = True
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(4)
        for b in role["bullets"]:
            bullet(b)
    if resume.get("projects"):
        heading("Projects")
        for project in resume["projects"]:
            p = doc.add_paragraph()
            p.add_run(project["name"]).bold = True
            if project.get("link"):
                p.add_run(f"  {project['link']}")
            p.paragraph_format.space_after = Pt(1)
            for b in project["bullets"]:
                bullet(b)
    heading("Education")
    for item in resume.get("education", []):
        para(f"{item['name']}    {item.get('dates', '')}".rstrip(), space_after=1)
    heading("Certifications")
    for cert in resume.get("certifications", []):
        bullet(cert)
    if resume.get("awards"):
        heading("Awards")
        for award in resume["awards"]:
            bullet(award)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


_LATIN1 = str.maketrans({"–": "-", "—": "-", "‘": "'", "’": "'", "“": '"', "”": '"', "•": "-", "…": "...", " ": " "})


def _pdf_text(text: str) -> str:
    """The built-in PDF fonts only cover Latin-1, so swap typographic characters for plain ones."""
    text = text.translate(_LATIN1)
    return re.sub(r"[^\x00-\xff]", "", text)


def to_pdf(resume: dict) -> bytes:
    pdf = FPDF(format="Letter")
    pdf.set_margins(18, 15, 18)
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def line(text, style="", size=10, align="L", gap=0.5):
        pdf.set_font("Helvetica", style, size)
        pdf.multi_cell(width, size * 0.48, _pdf_text(text), align=align, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)

    def heading(text):
        pdf.ln(2)
        line(text.upper(), "B", 11)
        y = pdf.get_y()
        pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
        pdf.ln(1.5)

    def bullet(text):
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(4, 4.8, "-")
        pdf.multi_cell(width - 4, 4.8, _pdf_text(text), align="L", new_x="LMARGIN", new_y="NEXT")

    contact = resume["contact"]
    line(contact["name"], "B", 18, "C", 1)
    line(resume["headline"], "", 11, "C", 0.5)
    line(contact_line(contact), "", 8.5, "C", 2)
    heading("Summary")
    line(resume["summary"])
    heading("Skills")
    for group in resume["skills"]:
        line(f"{group['name']}: {', '.join(group['items'])}", gap=0)
    heading("Experience")
    for role in resume["experience"]:
        pdf.ln(1)
        line(f"{role['title']}, {role['org']}  ({role['dates']})", "B", 10, gap=0)
        for b in role["bullets"]:
            bullet(b)
    if resume.get("projects"):
        heading("Projects")
        for project in resume["projects"]:
            line(f"{project['name']}  {project.get('link', '')}".rstrip(), "B", 10, gap=0)
            for b in project["bullets"]:
                bullet(b)
    heading("Education")
    for item in resume.get("education", []):
        line(f"{item['name']}  ({item.get('dates', '')})", gap=0)
    heading("Certifications")
    for cert in resume.get("certifications", []):
        bullet(cert)
    if resume.get("awards"):
        heading("Awards")
        for award in resume["awards"]:
            bullet(award)
    return bytes(pdf.output())
